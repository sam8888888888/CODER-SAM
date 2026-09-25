#!/usr/bin/env python3
# Membuat berkas uji pratinjau untuk e2e/ui.e2e.mjs (docx, xlsx, pptx).
#
# docx memakai python-docx, xlsx memakai openpyxl, pptx dibuat langsung sebagai paket OOXML (zip)
# karena tidak ada pustaka penulis pptx di mesin ini. Berkas hasilnya kecil dan statis, sehingga
# uji e2e tidak perlu pustaka Python.
#
# Jalankan dari akar repo dashboard:  python3 e2e/make-fixtures.py
import os
import zipfile

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(HERE, "fixtures")
os.makedirs(OUT, exist_ok=True)
TEXT_DOCX = "DOKUMEN UJI WAVE 11A"
TEXT_XLSX = "SEL-UJI-11A"
TEXT_PPTX = "SLIDE UJI WAVE 11A"


def make_docx(path):
    from docx import Document

    document = Document()
    document.add_heading(TEXT_DOCX, level=1)
    document.add_paragraph("Paragraf pertama untuk menguji penampil pratinjau berkas Word.")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Kolom A"
    table.cell(0, 1).text = "Kolom B"
    table.cell(1, 0).text = "1"
    table.cell(1, 1).text = "2"
    document.save(path)


def make_xlsx(path):
    from openpyxl import Workbook

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Uji"
    sheet["A1"] = TEXT_XLSX
    sheet["B1"] = 42
    sheet["A2"] = "kedua"
    sheet["B2"] = "baris"
    second = workbook.create_sheet("Kedua")
    second["A1"] = "LEMBAR-KEDUA"
    workbook.save(path)


A_NS = 'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"'
P_NS = 'xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main"'
R_NS = 'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"'
RELS_NS = 'xmlns="http://schemas.openxmlformats.org/package/2006/relationships"'
XML_HEAD = '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'

CONTENT_TYPES = XML_HEAD + '''<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">
<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>
<Default Extension="xml" ContentType="application/xml"/>
<Override PartName="/ppt/presentation.xml" ContentType="application/vnd.openxmlformats-officedocument.presentationml.presentation.main+xml"/>
<Override PartName="/ppt/slideMasters/slideMaster1.xml" ContentType="application/vnd.openxmlformats-officedocument.presentationml.slideMaster+xml"/>
<Override PartName="/ppt/slideLayouts/slideLayout1.xml" ContentType="application/vnd.openxmlformats-officedocument.presentationml.slideLayout+xml"/>
<Override PartName="/ppt/slides/slide1.xml" ContentType="application/vnd.openxmlformats-officedocument.presentationml.slide+xml"/>
<Override PartName="/ppt/theme/theme1.xml" ContentType="application/vnd.openxmlformats-officedocument.theme+xml"/>
</Types>'''

ROOT_RELS = XML_HEAD + '<Relationships ' + RELS_NS + '''>
<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="ppt/presentation.xml"/>
</Relationships>'''

DEFAULT_TEXT_STYLE = '''
  <a:lvl1PPr marL="342900" indent="0"><a:defRPr sz="1800"/></a:lvl1PPr>
  <a:lvl2PPr marL="685800" indent="0"><a:defRPr sz="1750"/></a:lvl2PPr>
  <a:lvl3PPr marL="1028700" indent="0"><a:defRPr sz="1700"/></a:lvl3PPr>
  <a:lvl4PPr marL="1371600" indent="0"><a:defRPr sz="1650"/></a:lvl4PPr>
  <a:lvl5PPr marL="1714500" indent="0"><a:defRPr sz="1600"/></a:lvl5PPr>
  <a:lvl6PPr marL="2057400" indent="0"><a:defRPr sz="1550"/></a:lvl6PPr>
  <a:lvl7PPr marL="2400300" indent="0"><a:defRPr sz="1500"/></a:lvl7PPr>
  <a:lvl8PPr marL="2743200" indent="0"><a:defRPr sz="1450"/></a:lvl8PPr>
  <a:lvl9PPr marL="3086100" indent="0"><a:defRPr sz="1400"/></a:lvl9PPr>
'''

PRESENTATION = XML_HEAD + '<p:presentation ' + A_NS + ' ' + R_NS + ' ' + P_NS + ''' saveSubsetFonts="1">
<p:sldMasterIdLst><p:sldMasterId id="2147483648" r:id="rId1"/></p:sldMasterIdLst>
<p:sldIdLst><p:sldId id="256" r:id="rId2"/></p:sldIdLst>
<p:sldSz cx="9144000" cy="6858000" type="screen4x3"/>
<p:notesSz cx="6858000" cy="9144000"/>
<p:defaultTextStyle>''' + DEFAULT_TEXT_STYLE + '''</p:defaultTextStyle>
</p:presentation>'''

PRESENTATION_RELS = XML_HEAD + '<Relationships ' + RELS_NS + '''>
<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/slideMaster" Target="slideMasters/slideMaster1.xml"/>
<Relationship Id="rId2" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/slide" Target="slides/slide1.xml"/>
<Relationship Id="rId3" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/theme" Target="theme/theme1.xml"/>
</Relationships>'''

SHAPE_TEXT = '''<p:sp>
<p:nvSpPr><p:cNvPr id="2" name="KotakUji"/><p:cNvSpPr/><p:nvPr/></p:nvSpPr>
<p:spPr><a:xfrm><a:off x="914400" y="914400"/><a:ext cx="7315200" cy="1828800"/></a:xfrm>
<a:prstGeom prst="rect"><a:avLst/></a:prstGeom></p:spPr>
<p:txBody><a:bodyPr/><a:lstStyle/>
<a:p><a:r><a:rPr lang="id-ID" sz="3200" b="1"/><a:t>''' + TEXT_PPTX + '''</a:t></a:r></a:p>
<a:p><a:r><a:rPr lang="id-ID" sz="1800"/><a:t>Baris kedua untuk uji penampil pratinjau.</a:t></a:r></a:p>
</p:txBody>
</p:sp>'''

SP_TREE_HEAD = '''<p:spTree>
<p:nvGrpSpPr><p:cNvPr id="1" name=""/><p:cNvGrpSpPr/><p:nvPr/></p:nvGrpSpPr>
<p:grpSpPr><a:xfrm><a:off x="0" y="0"/><a:ext cx="0" cy="0"/><a:chOff x="0" y="0"/><a:chExt cx="0" cy="0"/></a:xfrm></p:grpSpPr>'''

SLIDE1 = XML_HEAD + '<p:sld ' + A_NS + ' ' + R_NS + ' ' + P_NS + '''>
<p:cSld>''' + SP_TREE_HEAD + SHAPE_TEXT + '''</p:spTree></p:cSld>
<p:clrMapOvr><a:masterClrMapping/></p:clrMapOvr>
</p:sld>'''

SLIDE1_RELS = XML_HEAD + '<Relationships ' + RELS_NS + '''>
<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/slideLayout" Target="../slideLayouts/slideLayout1.xml"/>
</Relationships>'''

SLIDE_MASTER = XML_HEAD + '<p:sldMaster ' + A_NS + ' ' + R_NS + ' ' + P_NS + '''>
<p:cSld>''' + SP_TREE_HEAD + '''</p:spTree></p:cSld>
<p:clrMap bg1="lt1" tx1="dk1" bg2="lt2" tx2="dk2" accent1="accent1" accent2="accent2" accent3="accent3" accent4="accent4" accent5="accent5" accent6="accent6" hlink="hlink" folHlink="folHlink"/>
<p:sldLayoutIdLst><p:sldLayoutId id="2147483649" r:id="rId1"/></p:sldLayoutIdLst>
<p:txStyles>
<p:titleStyle><a:lvl1pPr><a:defRPr sz="4400"/></a:lvl1pPr></p:titleStyle>
<p:bodyStyle><a:lvl1pPr><a:defRPr sz="2000"/></a:lvl1pPr></p:bodyStyle>
<p:otherStyle><a:lvl1pPr><a:defRPr sz="1800"/></a:lvl1pPr></p:otherStyle>
</p:txStyles>
</p:sldMaster>'''

SLIDE_MASTER_RELS = XML_HEAD + '<Relationships ' + RELS_NS + '''>
<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/slideLayout" Target="../slideLayouts/slideLayout1.xml"/>
<Relationship Id="rId2" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/theme" Target="../theme/theme1.xml"/>
</Relationships>'''

SLIDE_LAYOUT = XML_HEAD + '<p:sldLayout ' + A_NS + ' ' + R_NS + ' ' + P_NS + ''' type="blank" preserve="1">
<p:cSld name="Kosong">''' + SP_TREE_HEAD + '''</p:spTree></p:cSld>
<p:clrMapOvr><a:masterClrMapping/></p:clrMapOvr>
</p:sldLayout>'''

SLIDE_LAYOUT_RELS = XML_HEAD + '<Relationships ' + RELS_NS + '''>
<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/slideMaster" Target="../slideMasters/slideMaster1.xml"/>
</Relationships>'''

THEME = XML_HEAD + '''<a:theme xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" name="Tema Uji">
<a:themeElements>
<a:clrScheme name="Uji">
<a:dk1><a:srgbClr val="000000"/></a:dk1><a:lt1><a:srgbClr val="FFFFFF"/></a:lt1>
<a:dk2><a:srgbClr val="1F2937"/></a:dk2><a:lt2><a:srgbClr val="F3F4F6"/></a:lt2>
<a:accent1><a:srgbClr val="2563EB"/></a:accent1><a:accent2><a:srgbClr val="DB2777"/></a:accent2>
<a:accent3><a:srgbClr val="059669"/></a:accent3><a:accent4><a:srgbClr val="D97706"/></a:accent4>
<a:accent5><a:srgbClr val="7C3AED"/></a:accent5><a:accent6><a:srgbClr val="0891B2"/></a:accent6>
<a:hlink><a:srgbClr val="0563C1"/></a:hlink><a:folHlink><a:srgbClr val="954F72"/></a:folHlink>
</a:clrScheme>
<a:fontScheme name="Uji">
<a:majorFont><a:latin typeface="Calibri Light"/><a:ea typeface=""/><a:cs typeface=""/></a:majorFont>
<a:minorFont><a:latin typeface="Calibri"/><a:ea typeface=""/><a:cs typeface=""/></a:minorFont>
</a:fontScheme>
<a:fmtScheme name="Uji">
<a:fillStyleLst><a:solidFill><a:schemeClr val="phClr"/></a:solidFill><a:solidFill><a:schemeClr val="phClr"/></a:solidFill><a:solidFill><a:schemeClr val="phClr"/></a:solidFill></a:fillStyleLst>
<a:lnStyleLst><a:ln w="6350"><a:solidFill><a:schemeClr val="phClr"/></a:solidFill></a:ln><a:ln w="12700"><a:solidFill><a:schemeClr val="phClr"/></a:solidFill></a:ln><a:ln w="19050"><a:solidFill><a:schemeClr val="phClr"/></a:solidFill></a:ln></a:lnStyleLst>
<a:effectStyleLst><a:effectStyle><a:effectLst/></a:effectStyle><a:effectStyle><a:effectLst/></a:effectStyle><a:effectStyle><a:effectLst/></a:effectStyle></a:effectStyleLst>
<a:bgFillStyleLst><a:solidFill><a:schemeClr val="phClr"/></a:solidFill><a:solidFill><a:schemeClr val="phClr"/></a:solidFill><a:solidFill><a:schemeClr val="phClr"/></a:solidFill></a:bgFillStyleLst>
</a:fmtScheme>
</a:themeElements>
<a:objectDefaults/><a:extraClrSchemeLst/>
</a:theme>'''


def make_pptx(path):
    parts = {
        "[Content_Types].xml": CONTENT_TYPES,
        "_rels/.rels": ROOT_RELS,
        "ppt/presentation.xml": PRESENTATION,
        "ppt/_rels/presentation.xml.rels": PRESENTATION_RELS,
        "ppt/slides/slide1.xml": SLIDE1,
        "ppt/slides/_rels/slide1.xml.rels": SLIDE1_RELS,
        "ppt/slideMasters/slideMaster1.xml": SLIDE_MASTER,
        "ppt/slideMasters/_rels/slideMaster1.xml.rels": SLIDE_MASTER_RELS,
        "ppt/slideLayouts/slideLayout1.xml": SLIDE_LAYOUT,
        "ppt/slideLayouts/_rels/slideLayout1.xml.rels": SLIDE_LAYOUT_RELS,
        "ppt/theme/theme1.xml": THEME,
    }
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as archive:
        for name, body in parts.items():
            archive.writestr(name, body)


MARKDOWN = '''# Catatan Uji Wave 11A

CATATAN-MARKDOWN-WAVE-11A

Berkas ini dipakai uji antarmuka untuk pratinjau markdown dan uji sunting artefak.

- butir 49: pratinjau markdown
- butir 48: sunting artefak dan riwayat revisi
'''


def make_markdown(path):
    with open(path, "w", encoding="utf-8") as handle:
        handle.write(MARKDOWN)


def main():
    make_docx(os.path.join(OUT, "sample.docx"))
    make_xlsx(os.path.join(OUT, "sample.xlsx"))
    make_pptx(os.path.join(OUT, "sample.pptx"))
    make_markdown(os.path.join(OUT, "catatan-uji.md"))
    for name in ("sample.docx", "sample.xlsx", "sample.pptx", "catatan-uji.md"):
        print(name, os.path.getsize(os.path.join(OUT, name)), "byte")


if __name__ == "__main__":
    main()
